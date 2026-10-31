"""Сверка «Получено от плательщиков» (HolliHop, expertonline.t8s.ru) с таблицей учеников (xlsx).

Вход: xlsx с листами «<ПРЕДМЕТ> ГК» и «<ПРЕДМЕТ> спецкурсы» + PDF со ссылками на группы.
Выход: docx с расхождениями. Вход на сайт человек делает сам в окне приложения (обычная страница входа),
дальше окно по очереди открывает страницы групп из PDF, программа читает их и ничего на сайте не меняет.

Без аргументов открывается окно. Консоль (окно сайта для входа всё равно откроется):
    python sverka.py --xlsx table.xlsx --pdf courses.pdf [--out report.docx]
"""
import argparse
import datetime as dt
import difflib
import hashlib
import html
import json
import os
import re
import ssl
import subprocess
import sys
import tempfile
import threading
import urllib.parse
import urllib.request
import warnings
import webbrowser
from dataclasses import dataclass, field
from decimal import Decimal
from pathlib import Path

import egespace
import shkolkovo

import docx
import openpyxl
import pypdf
import webview

import daily
from docx.enum.section import WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.shared import Pt, RGBColor
from openpyxl.utils import column_index_from_string, get_column_letter

MONTHS = [("сентяб", "Сентябрь"), ("октяб", "Октябрь"), ("ноябр", "Ноябрь"), ("декаб", "Декабрь"),
          ("январ", "Январь"), ("феврал", "Февраль"), ("март", "Март"), ("апрел", "Апрель"),
          ("май", "Май"), ("мая", "Май"), ("июн", "Июнь"), ("июл", "Июль"), ("август", "Август")]
SUBJECTS = {"общ": "ОБЩ", "ист": "ИСТ", "рус": "РУС", "мат": "МАТ"}
SUBJECT_WORDS = {"обществознан": "ОБЩ", "истори": "ИСТ", "русск": "РУС", "математик": "МАТ"}
ALIASES = {"с-п": "стартер-пак"}  # сокращение в названии группы на сайте → подпись в xlsx
LATIN_LOOKALIKES = str.maketrans("aceopxykmthb", "асеорхукмтнв")
SITE = "https://expertonline.t8s.ru"  # пароль уходит только сюда; ссылки PDF на другие адреса пропускаются
GROUP_PATH = re.compile(r"/Learner/(?:Group/)?(\d+)/?")
ZERO = Decimal(0)
try:
    from build_info import BUILD  # номер сборки пишет CI перед PyInstaller
except ImportError:
    BUILD = 0  # запуск из исходников — обновления не проверяем
RELEASES_API = "https://api.github.com/repos/Lukashev97/expert-sverka/releases/latest"


class LoginError(Exception):
    pass


# ---------- входные файлы ----------

def pdf_links(path, site=SITE):
    """Ссылки на группы сайта из PDF в порядке появления, по одной на группу; остальные — в skipped."""
    origin = urllib.parse.urlsplit(site)
    links, ids, skipped = [], set(), []
    for page in pypdf.PdfReader(path).pages:
        annots = page.get("/Annots")
        for a in (annots.get_object() if annots else []):
            act = a.get_object().get("/A")
            uri = act.get_object().get("/URI") if act else None
            if not uri:
                continue
            uri = str(uri).strip()
            u = urllib.parse.urlsplit(uri)
            m = GROUP_PATH.fullmatch(u.path)
            if m and not u.query and (u.scheme, u.netloc.casefold()) == (origin.scheme, origin.netloc.casefold()):
                if m[1] not in ids:  # /Learner/387 и /Learner/Group/387 — одна группа
                    ids.add(m[1])
                    links.append(uri)
            elif uri not in skipped:
                skipped.append(uri)
    return links, skipped


def norm(s):
    return re.sub(r"\s+", "", str(s).casefold().translate(LATIN_LOOKALIKES))


def month_of(token):
    return next((m for stem, m in MONTHS if token.startswith(stem)), None)


def subject_of(token):
    return SUBJECTS.get(token) or next((v for k, v in SUBJECT_WORDS.items() if token.startswith(k)), None)


@dataclass
class Entry:  # одна итоговая ячейка xlsx
    sheet: str
    subject: str
    month: str | None
    label: str  # нормализованная подпись, напр. «профи2»
    value: Decimal
    cell: str
    students: list | None = None  # [(ФИО, сумма, строка)] — кого складывает SUMIF ячейки; None — формула не SUMIF


SUMIF_RE = re.compile(r'=SUMIF\(([A-Z]+)(\d+):([A-Z]+)(\d+),(?:([A-Z]+)(\d+)|"([^"]*)"),([A-Z]+)(\d+):([A-Z]+)(\d+)\)')


def sumif_rows(ws, formula):
    """Ученики, которых складывает формула сводки =SUMIF(диапазон; критерий; суммы): [(ФИО, сумма, строка)]."""
    m = SUMIF_RE.fullmatch(str(formula or "").replace("$", "").replace(" ", ""))
    if not m:
        return None
    crit = ws[f"{m[5]}{m[6]}"].value if m[5] else m[7]
    cc, sc, shift = column_index_from_string(m[1]), column_index_from_string(m[8]), int(m[9]) - int(m[2])
    name_col = next((c for c in range(1, 8) if norm(ws.cell(1, c).value or "") == "фио"), 2)
    return [(ws.cell(r, name_col).value, _dec(ws.cell(r + shift, sc).value), r)
            for r in range(int(m[2]), min(int(m[4]), ws.max_row) + 1)
            if norm(ws.cell(r, cc).value or "") == norm(crit) and _is_num(ws.cell(r + shift, sc).value)]


def load_xlsx(path):
    """Итоговые ячейки из сводок наверху листов «… ГК» (месяц × тариф) и «… спецкурсы» (курс),
    замечания о строках сводки без суммы и где в таблице встречается каждый ученик (по фамилии и инициалу)."""
    entries, problems, people = [], [], {}
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = openpyxl.load_workbook(path, data_only=True)  # значения
        wbf = openpyxl.load_workbook(path)  # формулы сводок — по ним видно, каких учеников складывает ячейка
    for ws in wb.worksheets:
        tokens = re.findall(r"[\wё-]+", ws.title.casefold())
        subject = next(filter(None, map(subject_of, tokens)), None)
        if not subject:
            continue
        rows = [list(r) for r in ws.iter_rows(max_row=60, values_only=True)]
        width = max((len(r) for r in rows), default=0)
        rows = [r + [None] * (width - len(r)) for r in rows]
        if not rows:
            continue
        wf = wbf[ws.title]
        name_col = next((c for c in range(1, 8) if norm(ws.cell(1, c).value or "") == "фио"), 2)
        for r in range(2, ws.max_row + 1):
            name = ws.cell(r, name_col).value
            if isinstance(name, str) and len(name.split()) > 1:
                people.setdefault(name_key(name), []).append((ws.title, r, name))

        def amount(r, c):  # (сумма, ученики) ячейки r,c (0-based); без кэша значения — считаем SUMIF сами
            students = sumif_rows(ws, wf.cell(r + 1, c + 1).value)
            value = rows[r][c] if _is_num(rows[r][c]) else None
            if value is None and students is not None:
                value = sum((a for _, a, _ in students), ZERO)
            return value, students
        if "гк" in tokens:
            # Шапка: месяцы дважды — сначала оплаты, потом тарифы. Сводка «подпись / сумма под ней»
            # стоит в колонках второго набора месяцев.
            cols = {}
            for c, h in enumerate(rows[0]):
                m = month_of(norm(h)) if isinstance(h, str) else None
                if m:
                    cols.setdefault(m, []).append(c)
            for m, cs in cols.items():
                if len(cs) < 2:
                    continue
                c = cs[-1]
                for r in range(1, len(rows) - 1):
                    label = rows[r][c]
                    if not (isinstance(label, str) and label.strip()):
                        continue
                    value, students = amount(r + 1, c)
                    if value is not None:
                        lab = norm(label).removeprefix("гк")
                        entries.append(Entry(ws.title, subject, m, lab, _dec(value), f"{get_column_letter(c + 1)}{r + 2}", students))
        elif any(t.startswith("спецкурс") for t in tokens):
            # Строки «Количество | N | сумма | Курс»
            for r in range(1, len(rows)):
                a, c, d = rows[r][0], rows[r][2] if width > 2 else None, rows[r][3] if width > 3 else None
                if not (isinstance(a, str) and a.strip().casefold() == "количество"):
                    continue
                value, students = amount(r, 2)
                if value is None:
                    problems.append(f"«{ws.title}», строка {r + 1}: в C вместо суммы {f'«{c}»' if c is not None else 'пусто'}")
                elif isinstance(d, str) and norm(d) != "все":
                    entries.append(Entry(ws.title, subject, None, norm(d), _dec(value), f"C{r + 1}", students))
    return entries, problems, people


def _is_num(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def _dec(v):
    return Decimal(str(v)).quantize(Decimal("0.01"))


# ---------- сайт (только чтение) ----------

def group_path(url, site=SITE):
    """Путь страницы группы на нашем сайте; всё остальное читать отказываемся."""
    u, o = urllib.parse.urlsplit(url), urllib.parse.urlsplit(site)
    if (u.scheme, u.netloc.casefold()) != (o.scheme, o.netloc.casefold()) or not GROUP_PATH.fullmatch(u.path):
        raise ValueError(f"не страница группы {o.netloc}: {url}")
    return u.path


READ_PAGE_JS = """(() => {
  const cap = document.querySelector("x-caption");
  const td = Array.from(document.querySelectorAll("td")).find(t => t.textContent.trim() === "Получено от плательщиков:");
  return {path: location.pathname,
          html: "<title>" + document.title + "</title>" + (cap ? cap.outerHTML : "") + (td ? td.parentElement.outerHTML : "")};
})()"""
PAYMENTS = "/Payment/SearchLearnerPayments"
# Список оплат группы. В строках есть кнопки «изменить»/«удалить» — их не трогаем, читаем только текст таблицы.
READ_PAYMENTS_JS = """(() => {
  const t = Array.from(document.querySelectorAll("table")).find(t => t.querySelector("th") && t.textContent.includes("Плательщик"));
  const text = n => n.textContent.replace(/\\s+/g, " ").trim();  // по узлам: «Кузьмин И. Н.» и «заказчик: …» не слипаются
  const cells = tr => Array.from(tr.children).map(c => Array.from(c.childNodes).map(text).filter(Boolean).join(" "));
  const pages = Array.from(document.querySelectorAll("a.page-link"))
    .map(a => +(((a.getAttribute("href") || "").match(/[?&]Page=(\\d+)/) || [0, 0])[1]));
  return {path: location.pathname, search: location.search,
          head: t ? Array.from(t.querySelectorAll("th")).map(th => th.textContent.trim()) : [],
          rows: t ? Array.from(t.querySelectorAll("tr")).slice(1).map(cells) : [],
          pages: Math.max(0, ...pages) + 1,
          footer: ((document.body.textContent.match(/Всего найдено[^\\n]*/) || [""])[0]).trim()};
})()"""
SAFE_PATH = re.compile(r"/Learner/(?:Group/)?\d+/?|/Payment/SearchLearnerPayments\?LearnerId=\d+(?:&Page=\d+)?")
CLOSED = "Окно сайта закрыто — сверка прервана."


class WebSession:
    """Окно сайта: человек сам входит на обычной странице входа, дальше читаем страницы групп его сессией."""

    def __init__(self, site):
        self.site = site.rstrip("/")
        self.closed = threading.Event()
        # без ReturnUrl: после входа сайт уходит на свою главную, и первая группа — всегда новый адрес
        self.win = webview.create_window("Вход на сайт — войдите своим логином и паролем",
                                         f"{self.site}/Account/LogOn", width=1000, height=760)
        self.win.events.closed += self.closed.set

    def wait_login(self):
        host = urllib.parse.urlsplit(self.site).netloc.casefold()
        while not self.closed.wait(0.5):
            try:  # страница входа может грузиться долго — ждём, пока человек не войдёт или не закроет окно
                u = urllib.parse.urlsplit(self.win.get_current_url() or "")
                if u.netloc.casefold() == host and not u.path.casefold().startswith("/account/") \
                        and self.win.evaluate_js("document.readyState") == "complete":
                    return
            except Exception:
                pass
        raise LoginError("Окно сайта закрыто — сверка отменена.")

    def _url(self):
        try:
            return self.win.get_current_url() or ""
        except Exception:
            return ""

    def _open(self, path, again):
        if not SAFE_PATH.fullmatch(path):  # открываем только страницы групп и списки их оплат
            raise ValueError(f"не читаем {path}")
        u = urllib.parse.urlsplit(self._url())
        if again and (u.path + (f"?{u.query}" if u.query else "")).rstrip("/") == path.rstrip("/"):
            # повтор: WebView2 молча игнорирует load_url на уже открытую страницу — пусть перезагрузится сама
            self.win.evaluate_js("setTimeout(() => location.reload(), 50)")
            self.win.events.loaded.clear()
        else:
            self.win.load_url(self.site + path)
        for _ in range(240):  # до 2 минут, но сразу выходим, если окно закрыли
            if self.win.events.loaded.wait(0.5):
                return
            if self.closed.is_set():
                raise LoginError(CLOSED)
        raise TimeoutError("страница не загрузилась за 2 минуты")

    def _read(self, path, js, ok):
        """Окно само открывает страницу (как клик по ссылке) и отдаёт то, что вернул js, если ok(результат)."""
        err = None
        for attempt in range(2):
            if self.closed.is_set():
                raise LoginError(CLOSED)
            self._open(path, again=attempt > 0)
            if "/account/logon" in self._url().casefold():
                raise LoginError("Сайт попросил войти заново (сессия закончилась).")
            try:
                r = self.win.evaluate_js(js)
            except Exception as e:  # страница успела перезагрузиться сама — пробуем ещё раз
                err = e
                continue
            if isinstance(r, dict) and ok(r):
                return r
            err = f"открылась не та страница: {r.get('path') if isinstance(r, dict) else r}"
        raise ValueError(f"страница не прочиталась: {err}")

    def get(self, url):
        """Страница группы: заголовок и строка статистики."""
        path = group_path(url, self.site)
        want = GROUP_PATH.fullmatch(path)[1]
        r = self._read(path, READ_PAGE_JS, lambda r: (m := GROUP_PATH.fullmatch(r.get("path", ""))) and m[1] == want)
        return r["html"]

    def payments(self, url):
        """Оплаты группы со всех страниц списка «Оплаты»: [(плательщик, сумма)] и строка «Всего найдено …»."""
        gid = GROUP_PATH.fullmatch(group_path(url, self.site))[1]
        out, footer, page, pages = [], "", 0, 1
        while page < min(pages, 300):
            q = f"LearnerId={gid}" + (f"&Page={page}" if page else "")
            r = self._read(f"{PAYMENTS}?{q}", READ_PAYMENTS_JS,
                           lambda r: r.get("path") == PAYMENTS and f"LearnerId={gid}" in r.get("search", ""))
            out += payment_rows(r["head"], r["rows"])
            pages, footer, page = r["pages"], r["footer"], page + 1
        return out, footer

    def close(self):
        if not self.closed.is_set():
            self.win.destroy()


def login_window(pdf, site=SITE):
    links, _ = pdf_links(pdf, site)
    if not links:
        raise ValueError(f"В PDF нет ссылок на группы {site}/Learner/Group/<номер>.")
    return WebSession(site)


def parse_group_page(text):
    """(название группы, «Получено от плательщиков») со страницы группы."""
    m = re.search(r"<x-caption[^>]*>\s*([^<]*?)\s*<", text) or re.search(r"<title>\s*(?:Группа\s+)?([^<]*?)\s*</title>", text)
    name = html.unescape(m.group(1)).strip() if m else ""
    m = re.search(r"Получено от плательщиков:\s*</td>\s*<td[^>]*>(.*?)</td>", text, re.S)
    if not m:
        raise ValueError("на странице нет «Получено от плательщиков»")
    return name, parse_money(html.unescape(re.sub(r"<[^>]+>", "", m.group(1))))


def payment_rows(head, rows):
    """Строки таблицы оплат → [(плательщик, сумма)]; неоплаченные счета не считаются."""
    idx = {h.strip(): i for i, h in enumerate(head)}
    ip, ia, ist = idx.get("Плательщик"), idx.get("Размер оплаты"), idx.get("Статус")
    if ip is None or ia is None:
        raise ValueError("на странице оплат нет колонок «Плательщик» / «Размер оплаты»")
    return [(re.sub(r"\s*заказчик:\s*(.+)", r" (заказчик: \1)", cells[ip]), parse_money(cells[ia])) for cells in rows
            if len(cells) > max(ip, ia) and not (ist is not None and len(cells) > ist and "не оплачен" in cells[ist].casefold())]


def parse_money(s):
    m = re.search(r"-?\d+(?:,\d+)?", re.sub(r"\s", "", s).replace("−", "-"))
    if not m:
        raise ValueError(f"не число: {s!r}")
    return Decimal(m.group().replace(",", "."))


# ---------- сопоставление ----------

def parse_group_name(name):
    """«26/27 октябрь ГК ОБЩ ПРОФИ 2» → ('ОБЩ', 'Октябрь', 'профи2')."""
    s = re.sub(r"\b\d{2}/\d{2}\b", " ", name.casefold().translate(LATIN_LOOKALIKES))
    subject = next(filter(None, map(subject_of, re.findall(r"[\wё+-]+", s))), None)
    head = re.split(r"[:/]", s, maxsplit=1)[0]  # «на пересдачу: математика / базовый минимум» → «на пересдачу»
    month, label = None, []
    for t in re.findall(r"[\wё+-]+", head):
        if t == "гк" or subject_of(t):
            continue
        if not month and month_of(t):
            month = month_of(t)
            continue
        label.append(t)
    label = "".join(label)
    return subject, month, ALIASES.get(label, label)


@dataclass
class Group:
    url: str
    name: str = ""
    site: Decimal | None = None
    error: str | None = None


@dataclass
class Row:  # одна строка отчёта: группа (или несколько групп на одну ячейку xlsx)
    groups: list
    status: str  # ключ из STATUS
    xlsx: Decimal | None = None
    cells: list = field(default_factory=list)
    note: str = ""
    hits: list = field(default_factory=list)  # ячейки xlsx (Entry), с которыми сравнивали
    students: list = field(default_factory=list)  # [(ученик на сайте, ученик в xlsx, сайт, xlsx)] — только расходящиеся
    students_note: str = ""
    typos: list = field(default_factory=list)  # [(на сайте, в xlsx)] — суммы равны, ФИО отличается опечаткой

    @property
    def site(self):
        vals = [g.site for g in self.groups if g.site is not None]
        return sum(vals, ZERO) if vals else None

    @property
    def diff(self):
        return None if self.site is None or self.xlsx is None else self.site - self.xlsx


def match(groups, entries):
    rows, by_cells = [], {}
    for g in groups:
        if g.error:
            rows.append(Row([g], "error", note=g.error))
            continue
        subject, month, label = parse_group_name(g.name)
        hits = [e for e in entries if e.label == label
                and (subject is None or e.subject == subject)
                and (month is None or e.month == month)]
        sheets_by_subject = {}
        for e in hits:
            sheets_by_subject.setdefault(e.subject, set()).add(e.sheet)
        if not label or not hits:
            where = ", ".join(filter(None, [subject, month])) or "все листы"
            rows.append(Row([g], "missing" if g.site else "empty", note=f"в сводке xlsx ({where}) нет строки для этой группы"))
        elif any(len(s) > 1 for s in sheets_by_subject.values()):
            sheets = ", ".join(f"«{s}»" for ss in sheets_by_subject.values() for s in sorted(ss))
            rows.append(Row([g], "ambiguous", note=f"подходит несколько листов: {sheets}"))
        else:
            key = tuple(sorted((e.sheet, e.cell) for e in hits))
            if key in by_cells:  # несколько групп на сайте → одна ячейка xlsx: сравниваем сумму
                by_cells[key].groups.append(g)
                continue
            row = Row([g], "", xlsx=sum((e.value for e in hits), ZERO), cells=[f"«{s}»!{c}" for s, c in key], hits=hits)
            by_cells[key] = row
            rows.append(row)
    for row in by_cells.values():
        row.status = "ok" if abs(row.diff) < Decimal("0.01") else "diff"
        if len(row.groups) > 1:
            row.note = "на сайте несколько групп, в xlsx одна сумма — сравнивается их сумма"
    return rows


def _name_tokens(name):
    return re.findall(r"[a-zа-я]+", str(name or "").casefold().replace("ё", "е").translate(LATIN_LOOKALIKES))


def name_key(name, swap=False):
    """«Иванова Анна Сергеевна» и «Иванова А. С.» → ('иванова', 'а'): на сайте плательщик с инициалами.
    swap — для имён, записанных в xlsx как «Имя Фамилия Отчество»."""
    t = _name_tokens(name)
    if swap and len(t) > 1:
        t[0], t[1] = t[1], t[0]
    return (t[0], t[1][0]) if len(t) > 1 else (t[0] if t else "", "")


def same_person(a, b):
    """Та же фамилия, а имя и отчество совпадают — целиком или по инициалу, если где-то записаны инициалами."""
    ta, tb = _name_tokens(a), _name_tokens(b)
    return bool(ta and tb) and ta[0] == tb[0] and all(
        x[0] == y[0] and (len(x) == 1 or len(y) == 1 or x == y) for x, y in zip(ta[1:], tb[1:]))


def student_diff(row, payments, people=None, skip_sheets=()):
    """Кто из учеников даёт расхождение: оплаты групп на сайте против учеников, которых складывает ячейка xlsx."""
    site, footers = [], []
    for g in row.groups:
        rows, footer = payments(g.url)
        site += rows
        footers.append(footer)
    if any(e.students is None for e in row.hits):  # иначе все плательщики сайта выглядели бы расхождением
        row.students_note = "в xlsx сумма не формулой SUMIF — по ученикам не сравнить"
        return
    xl = [(name, amount) for e in row.hits for name, amount, _ in (e.students or [])]
    counted = {(e.sheet, r) for e in row.hits for _, _, r in (e.students or [])}
    notes = []
    paid = sum((a for _, a in site), ZERO)
    if row.site is not None and abs(paid - row.site) >= Decimal("0.01"):
        notes.append(f"на сайте в списке оплат {fmt(paid)}, а «Получено от плательщиков» {fmt(row.site)}")
    site_keys = {name_key(n) for n, _ in site}

    def key(name, side):
        k = name_key(name)
        if side and k not in site_keys and name_key(name, swap=True) in site_keys:
            return name_key(name, swap=True)  # в xlsx «Имя Фамилия», на сайте «Фамилия И.»
        return k

    by = {}
    for side, items in ((0, site), (1, xl)):
        for name, amount in items:
            k = by.setdefault(key(name, side), [set(), set(), ZERO, ZERO])
            k[side].add(str(name).strip())
            k[2 + side] += amount

    def elsewhere(k, names):  # где ещё в таблице этот ученик — частая причина: записан не в ту группу/лист
        spots = [f"«{sh.strip()}» стр. {r}" for sh, r, full in (people or {}).get(k, [])
                 if (sh, r) not in counted and sh not in skip_sheets and any(same_person(full, n) for n in names)]
        return ", ".join(spots[:6]) + (" …" if len(spots) > 6 else "")

    differ = {k: v for k, v in by.items() if abs(v[2] - v[3]) >= Decimal("0.01")}
    # «только на сайте» + «только в xlsx» на ту же сумму и с почти той же фамилией — опечатка в ФИО, а не деньги
    for ks in [k for k, v in differ.items() if not v[1]]:
        kx = next((k for k, v in differ.items() if not v[0] and k[1] == ks[1] and v[3] == differ[ks][2]
                   and difflib.SequenceMatcher(None, k[0], ks[0]).ratio() >= 0.8), None)
        if kx:
            row.typos.append((", ".join(sorted(differ[ks][0])), ", ".join(sorted(differ[kx][1]))))
            del differ[ks], differ[kx]
    row.students = sorted((", ".join(sorted(s)) or "—", ", ".join(sorted(x)) or "—", a, b, elsewhere(k, s | x))
                          for k, (s, x, a, b) in differ.items())
    row.students_note = "; ".join(notes + [f for f in footers if f])


def reconcile(xlsx, pdf, fetch, progress=lambda i, n, msg: None, site=SITE, payments=None):
    links, skipped = pdf_links(pdf, site)
    entries, problems, people = load_xlsx(xlsx)
    groups = []
    for i, url in enumerate(links, 1):
        g = Group(url)
        try:
            g.name, g.site = parse_group_page(fetch(url))
        except LoginError:
            raise
        except Exception as e:  # одна битая страница не должна ронять всю сверку
            g.error = f"{type(e).__name__}: {e}"
        progress(i, len(links), f"{g.name or url}: {fmt(g.site) if g.error is None else g.error}")
        groups.append(g)
    rows = match(groups, entries)
    if payments:  # по группам с расхождением — кто из учеников его даёт
        # листы со сводкой, ни одна ячейка которой не совпала с сайтом, — прошлый год: в «где ещё» не смотрим
        skip = {e.sheet for e in entries} - {e.sheet for r in rows for e in r.hits}
        # все группы с деньгами: сумма может совпасть, а ошибки двух учеников — погасить друг друга
        todo = [r for r in rows if r.status in PROBLEM or (r.status == "ok" and (r.site or r.xlsx))]
        for i, row in enumerate(todo, 1):
            progress(i, len(todo), f"Оплаты по ученикам: {', '.join(g.name for g in row.groups)}")
            try:
                student_diff(row, payments, people, skip)
            except LoginError:
                raise
            except Exception as e:
                row.students_note = f"не удалось разобрать по ученикам: {type(e).__name__}: {e}"
            if row.status == "ok" and row.students:
                row.status = "students"
    return rows, links, skipped, problems


# ---------- отчёт ----------

STATUS = {"ok": "совпадает", "diff": "РАСХОЖДЕНИЕ", "missing": "РАСХОЖДЕНИЕ: нет в xlsx",
          "students": "РАСХОЖДЕНИЕ по ученикам", "empty": "нет в xlsx, на сайте 0", "ambiguous": "неоднозначно",
          "error": "ошибка загрузки"}
PROBLEM = ("diff", "missing", "students")
RED = RGBColor(0xC0, 0x00, 0x00)


def fmt(d):
    if d is None:
        return "—"
    return f"{d:,.2f}".replace(",", " ").replace(".", ",")


def _link(paragraph, url, text):
    part = paragraph.part
    r_id = part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    for tag, attr in (("w:color", "0563C1"), ("w:u", "single")):
        el = OxmlElement(tag)
        el.set(qn("w:val"), attr)
        rpr.append(el)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    run.append(t)
    h.append(run)
    paragraph._p.append(h)


def _table(doc, header, rows, red=lambda row: False):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    for cell, text in zip(t.rows[0].cells, header):
        cell.paragraphs[0].add_run(text).bold = True
    for row in rows:
        cells = t.add_row().cells
        for cell, value in zip(cells, row):
            p = cell.paragraphs[0]
            if isinstance(value, list):  # [(url, text), …] — ссылки на группы
                for j, (url, text) in enumerate(value):
                    if j:
                        p = cell.add_paragraph()
                    _link(p, url, text)
            else:
                run = p.add_run(str(value))
                if red(row):
                    run.font.color.rgb = RED
    return t


def write_report(path, rows, links, skipped, problems, xlsx, pdf):
    doc = docx.Document()
    sec = doc.sections[0]
    sec.orientation, sec.page_width, sec.page_height = WD_ORIENT.LANDSCAPE, sec.page_height, sec.page_width
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Pt(36))
    doc.styles["Normal"].font.size = Pt(9)

    doc.add_heading("Сверка «Получено от плательщиков»: сайт ↔ xlsx", level=1)
    doc.add_paragraph(f"Сформировано: {dt.datetime.now():%d.%m.%Y %H:%M}\n"
                      f"Таблица: {Path(xlsx).name}\nPDF: {Path(pdf).name} — ссылок на группы: {len(links)}")
    doc.add_paragraph(summary(rows))

    def cells_cell(r):
        c = r.cells
        return ", ".join(c) if len(c) <= 3 else f"{c[0]} … {c[-1]} (сумма {len(c)} ячеек)"

    def groups_cell(r):
        return [(g.url, g.name or g.url) for g in r.groups]

    def site_cell(r):
        if len(r.groups) == 1:
            return fmt(r.site)
        return " + ".join(fmt(g.site) for g in r.groups) + f" = {fmt(r.site)}"

    def students_cell(r):  # кто именно: «Петров В. Г.: сайт 45 000,00 / xlsx 45 105,00»
        lines = [f"{s if s != '—' else x}: сайт {fmt(a)} / xlsx {fmt(b)}" for s, x, a, b, _ in r.students[:8]]
        if len(r.students) > 8:
            lines.append(f"… ещё {len(r.students) - 8} — в разделе «Расхождения по ученикам»")
        return "\n".join(lines) or r.students_note

    head = ["Группа на сайте", "Сайт, руб.", "xlsx, руб.", "Разница (сайт − xlsx)", "Ученики с расхождением", "Ячейка xlsx", "Примечание"]
    doc.add_heading("Расхождения", level=2)
    diffs = [r for r in rows if r.status in PROBLEM]
    if diffs:
        _table(doc, head, [[groups_cell(r), site_cell(r), fmt(r.xlsx), fmt(r.diff), students_cell(r), cells_cell(r),
                            r.note or (STATUS[r.status] if r.status == "students" else "")] for r in diffs])
    else:
        doc.add_paragraph("Расхождений не найдено.")

    detailed = [r for r in diffs if r.students]
    if detailed:
        doc.add_heading("Расхождения по ученикам", level=2)
        for r in detailed:
            doc.add_heading(f"{' + '.join(g.name or g.url for g in r.groups)}: сайт {fmt(r.site)}, xlsx {fmt(r.xlsx)}", level=3)
            if r.students:
                _table(doc, ["Ученик на сайте", "Ученик в xlsx", "Сайт, руб.", "xlsx, руб.", "Разница (сайт − xlsx)", "Где ещё в xlsx"],
                       [[s, x, fmt(a), fmt(b), fmt(a - b), w] for s, x, a, b, w in r.students])
            if r.students_note:
                doc.add_paragraph(r.students_note)

    other = [r for r in rows if r.status in ("ambiguous", "error")]
    typos = [(g, s, x) for r in rows for s, x in r.typos for g in [" + ".join(g.name for g in r.groups)]]
    if typos:
        doc.add_heading("Возможные опечатки в ФИО (суммы совпадают)", level=2)
        _table(doc, ["Группа на сайте", "На сайте", "В xlsx"], [list(t) for t in typos])
    if other:
        doc.add_heading("Не удалось сверить", level=2)
        _table(doc, ["Группа на сайте", "Сайт, руб.", "Статус", "Причина"],
               [[groups_cell(r), site_cell(r), STATUS[r.status], r.note] for r in other])
    if problems:
        doc.add_heading("Замечания по сводке xlsx", level=2)
        for p in problems:
            doc.add_paragraph(p, style="List Bullet")
    if skipped:
        doc.add_heading("Пропущенные ссылки PDF (не группы или не " + urllib.parse.urlsplit(SITE).netloc + ")", level=2)
        for u in skipped:
            doc.add_paragraph(u, style="List Bullet")

    doc.add_heading("Все группы", level=2)
    _table(doc, ["Группа на сайте", "Сайт, руб.", "xlsx, руб.", "Разница", "Статус", "Ячейка xlsx"],
           [[groups_cell(r), site_cell(r), fmt(r.xlsx), fmt(r.diff), STATUS[r.status], cells_cell(r)] for r in rows],
           red=lambda row: row[4] in (STATUS[s] for s in PROBLEM))
    doc.save(path)


def run(xlsx, pdf, web, out=None, progress=lambda i, n, msg: None, site=SITE):
    progress(0, 0, "Войдите на сайт в открывшемся окне — логин и пароль вводите сами.")
    web.wait_login()
    progress(0, 0, "Вход выполнен, читаю страницы групп…")
    rows, links, skipped, problems = reconcile(xlsx, pdf, web.get, progress, site, payments=web.payments)
    out = out or str(Path(xlsx).with_name(f"Расхождения {dt.datetime.now():%Y-%m-%d %H-%M}.docx"))
    write_report(out, rows, links, skipped, problems, xlsx, pdf)
    return out, rows


def summary(rows):
    n = {s: sum(len(r.groups) for r in rows if r.status == s) for s in STATUS}
    return (f"Групп: {sum(n.values())}. Совпадает: {n['ok']}. Расхождений: {n['diff'] + n['missing'] + n['students']}"
            f" (из них нет в xlsx: {n['missing']}; сумма сходится, но не по ученикам: {n['students']})."
            f" Нет в xlsx при нуле на сайте: {n['empty']}."
            f" Неоднозначно: {n['ambiguous']}. Ошибок загрузки: {n['error']}.")


# ---------- обновление ----------

def _tls():
    # Собранное .app ищет сертификаты там, где они лежали у Python на машине сборки; на macOS берём системные.
    if sys.platform == "darwin" and os.path.exists("/etc/ssl/cert.pem"):
        return ssl.create_default_context(cafile="/etc/ssl/cert.pem")
    return ssl.create_default_context()


def _http(url, timeout):
    return urllib.request.urlopen(urllib.request.Request(url, headers={"User-Agent": "Sverka"}),
                                  timeout=timeout, context=_tls())


def asset_name():
    return "Sverka.exe" if sys.platform == "win32" else "Sverka-macos.zip"


def find_update(release, build=BUILD):
    """Сборка новее нашей из ответа GitHub releases/latest — или None."""
    m = re.fullmatch(r"build-(\d+)", str(release.get("tag_name", "")))
    asset = next((a for a in release.get("assets") or [] if a.get("name") == asset_name()), None)
    if not m or int(m[1]) <= build or not asset:
        return None
    return {"build": int(m[1]), "page": release.get("html_url", ""), "asset": asset}


def check_update(api=RELEASES_API, build=BUILD):
    if not build:
        return None
    with _http(api, 5) as r:
        return find_update(json.load(r), build)


def download(asset, dest, progress=lambda done, total: None):
    """Скачивает файл релиза и сверяет размер и SHA-256 (GitHub отдаёт его в поле digest)."""
    h, done, total = hashlib.sha256(), 0, asset["size"]
    with _http(asset["browser_download_url"], 60) as r, open(dest, "wb") as f:
        while chunk := r.read(1 << 16):
            f.write(chunk)
            h.update(chunk)
            done += len(chunk)
            if done == total or done % (1 << 20) < len(chunk):  # примерно раз в мегабайт
                progress(done, total)
    digest = asset.get("digest") or ""
    if done != total or (digest.startswith("sha256:") and digest[7:] != h.hexdigest()):
        os.remove(dest)
        raise ValueError("скачанный файл повреждён — попробуйте ещё раз")


# Запущенный .exe нельзя перезаписать: подменяет cmd после нашего выхода и запускает новую версию.
# Пути — в переменных окружения, чтобы кириллица в пути не ломалась о кодировку .cmd-файла.
SWAP_CMD = """@echo off
set n=0
:retry
move /y "%SVERKA_NEW%" "%SVERKA_EXE%" >nul 2>&1 && goto run
set /a n+=1
if %n% geq 120 goto end
ping -n 2 127.0.0.1 >nul
goto retry
:run
start "" "%SVERKA_EXE%" %SVERKA_ARGS%
:end
del "%~f0"
"""


def swap_after_exit(new, exe, args=""):
    script = Path(tempfile.gettempdir()) / f"sverka-update-{os.getpid()}.cmd"
    script.write_text(SWAP_CMD.replace("\n", "\r\n"), encoding="ascii")
    # PYINSTALLER_RESET_ENVIRONMENT: иначе новый .exe примет себя за дочерний процесс старого и будет искать
    # Python в его временной папке, которой уже нет («Failed to load Python DLL»)
    env = dict(os.environ, PYINSTALLER_RESET_ENVIRONMENT="1", SVERKA_NEW=str(new), SVERKA_EXE=str(exe), SVERKA_ARGS=args)
    subprocess.Popen(["cmd", "/c", str(script)], env=env, cwd=str(Path(exe).parent),
                     stdin=subprocess.DEVNULL, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL,
                     creationflags=subprocess.CREATE_NO_WINDOW)


def make_report(xlsx, first, last=None, out=None, say=print):
    """Ежедневные отчёты руководителю за день или диапазон — один docx рядом с таблицей."""
    book = daily.Book(xlsx)
    span = daily.days(first, last)
    if book.last_log_day and span[-1] > book.last_log_day:
        say(f"Внимание: журнал продаж в «{daily.ZP}» заполнен по {book.last_log_day:%d.%m.%Y} — "
            "у более поздних дней оплаты будут нулевыми.")
    texts = [book.report(d) for d in span]
    name = f"Отчет {span[0]:%d.%m.%Y}" + (f"–{span[-1]:%d.%m.%Y}" if len(span) > 1 else "") + ".docx"
    out = out or str(Path(xlsx).with_name(name))
    daily.write_docx(texts, out)
    return out


def open_file(path):
    if sys.platform == "win32":
        os.startfile(path)
    else:
        subprocess.run(["open" if sys.platform == "darwin" else "xdg-open", path], check=False)


# ---------- окно ----------

APP_HTML = """<!doctype html><meta charset="utf-8"><style>
body { font: 14px -apple-system, "Segoe UI", sans-serif; margin: 16px; background: #f4f4f4; color: #222 }
.row { display: flex; gap: 10px; align-items: center; margin: 8px 0 }
.row span { color: #555; overflow: hidden; text-overflow: ellipsis; white-space: nowrap }
button { padding: 6px 14px; font: inherit }
progress { width: 100%; margin: 10px 0 }
pre { background: #fff; border: 1px solid #ccc; height: 280px; overflow: auto; padding: 8px; margin: 0; white-space: pre-wrap }
</style>
<div class="row" id="upd" style="display: none"><b id="updtext"></b><button id="updbtn" onclick="doUpdate()"></button></div>
<div class="row"><button onclick="pick('xlsx')">Таблица xlsx…</button><span id="xlsx">не выбрана</span></div>
<div class="row"><button onclick="pick('pdf')">PDF со ссылками…</button><span id="pdf">не выбран</span></div>
<div class="row"><button id="go" onclick="go()">Сверить</button>
  <span>Откроется окно сайта — войдите в нём сами, дальше программа только читает страницы групп.</span></div>
<div class="row"><button id="rep" onclick="report()">Отчет</button>
  <input type="date" id="dfrom"> — <input type="date" id="dto">
  <span>одна дата или диапазон; нужна только таблица xlsx</span></div>
<hr>
<div class="row"><b>egespace.ru</b>
  <input id="eslogin" placeholder="логин" size="12"><input id="espass" type="password" placeholder="пароль" size="12"></div>
<div class="row"><button id="bk" onclick="backup()">Бэкап предметов</button>
  <input id="esids" placeholder="ид курсов через запятую" size="26">
  <span>пустое поле — все предметы</span></div>
<div class="row"><button id="bkdir" onclick="pickOut()">Папка для копии…</button><span id="esout"></span></div>
<div class="row"><button id="rs" onclick="restore()">Восстановить из бэкапа…</button>
  <span>папка с датой внутри egespace_backup; предметы создадутся заново</span></div>
<hr>
<div class="row"><b>Школково</b><span>логин и пароль берутся из полей выше — у egespace.ru и worldofege.ru общая учётка</span></div>
<div class="row"><button id="shq" onclick="addQuestions()">Задачи в тест</button>
  <input id="shsrc" placeholder="тема Школково: https://3.shkolkovo.online/catalog/13539" size="34">
  <input id="shdst" placeholder="тест: https://egespace.ru/extendedTests/…" size="30"></div>
<div class="row"><span style="width:7em"></span><input id="shnums" placeholder="номера: 1-15" size="12">
  <input id="shpts" type="number" min="1" value="1" size="3" title="баллов за задание">
  <select id="shtype">
    <option value="directComparison">прямое сравнение</option>
    <option value="set">символы в любом порядке</option>
    <option value="sequence">символы по порядку</option>
  </select>
  <span>пустые номера — вся тема</span></div>
<div class="row"><button id="shl" onclick="uploadLessons()">Выгрузить сборник Дощинского</button>
  <input id="shtasks" placeholder="задания ЕГЭ: 1-26" size="12">
  <span>уроки и тесты в курсе 35, раздел «Задания»; пусто — все №1–26</span></div>
<progress id="bar" value="0" max="1"></progress>
<pre id="log"></pre>
<script>
const files = {};
async function pick(kind) {
  const path = await pywebview.api.pick(kind);
  if (path) { files[kind] = path; document.getElementById(kind).textContent = path; }
}
function busy(on) { for (const id of ["go", "rep", "bk", "rs", "shq", "shl"]) document.getElementById(id).disabled = on; }
function esCreds() {
  const login = document.getElementById("eslogin").value.trim();
  const pass = document.getElementById("espass").value;
  if (!login || !pass) { say("Укажите логин и пароль egespace.ru."); return null; }
  return [login, pass];
}
async function pickOut() {
  const path = await pywebview.api.pick("folder");
  if (path) { files.esout = path; document.getElementById("esout").textContent = path; }
}
function backup() {
  const c = esCreds(); if (!c) return;
  busy(true);
  document.getElementById("log").textContent = "";
  pywebview.api.backup(c[0], c[1], document.getElementById("esids").value, files.esout || "");
}
function val(id) { return document.getElementById(id).value.trim(); }
function addQuestions() {
  const c = esCreds(); if (!c) return;
  if (!val("shsrc") || !val("shdst")) { say("Укажите ссылку на тему Школково и на тест."); return; }
  busy(true);
  document.getElementById("log").textContent = "";
  pywebview.api.add_questions(c[0], c[1], val("shsrc"), val("shdst"), val("shnums"),
                              Number(val("shpts")) || 1, val("shtype"));
}
function uploadLessons() {
  const c = esCreds(); if (!c) return;
  busy(true);
  document.getElementById("log").textContent = "";
  pywebview.api.upload_lessons(c[0], c[1], val("shtasks"));
}
async function restore() {
  const c = esCreds(); if (!c) return;
  const folder = await pywebview.api.pick("folder");
  if (!folder) return;
  busy(true);
  document.getElementById("log").textContent = "";
  pywebview.api.restore(c[0], c[1], folder);
}
function go() {
  if (!files.xlsx || !files.pdf) { say("Выберите таблицу xlsx и PDF."); return; }
  busy(true);
  document.getElementById("log").textContent = "";
  pywebview.api.start(files.xlsx, files.pdf);
}
function report() {
  const first = document.getElementById("dfrom").value, last = document.getElementById("dto").value;
  if (!files.xlsx) { say("Выберите таблицу xlsx."); return; }
  if (!first) { say("Выберите дату отчета."); return; }
  busy(true);
  document.getElementById("log").textContent = "";
  pywebview.api.report(files.xlsx, first, last || first);
}
{ const t = new Date(); document.getElementById("dfrom").value =
    new Date(t.getTime() - t.getTimezoneOffset() * 60000).toISOString().slice(0, 10); }
function say(msg) { const l = document.getElementById("log"); l.textContent += msg + "\\n"; l.scrollTop = l.scrollHeight; }
function progress(i, n) { const b = document.getElementById("bar"); b.max = Math.max(n, 1); b.value = i; }
function done() { busy(false); }
async function checkUpdate() {
  document.getElementById("esout").textContent = await pywebview.api.backup_dir();
  const u = await pywebview.api.check_update();
  if (!u) return;
  document.getElementById("updtext").textContent = "Доступна новая версия — сборка " + u.build + ".";
  document.getElementById("updbtn").textContent = u.auto ? "Обновить" : "Скачать";
  document.getElementById("upd").style.display = "flex";
}
function doUpdate() {
  if (document.getElementById("go").disabled) { say("Дождитесь окончания работы."); return; }
  document.getElementById("updbtn").disabled = true;
  pywebview.api.update();
}
if (window.pywebview && window.pywebview.api && window.pywebview.api.check_update) checkUpdate();
else window.addEventListener("pywebviewready", checkUpdate);
</script>"""


class Api:  # методы доступны странице приложения как pywebview.api.*
    _win = None

    def pick(self, kind):
        if kind == "folder":
            paths = self._win.create_file_dialog(webview.FileDialog.FOLDER)
            return paths[0] if paths else None
        types = ("Excel (*.xlsx)",) if kind == "xlsx" else ("PDF (*.pdf)",)
        paths = self._win.create_file_dialog(webview.FileDialog.OPEN, file_types=types)
        return paths[0] if paths else None

    def backup(self, login, password, ids, out=""):
        threading.Thread(target=self._egespace, args=("backup", login, password, ids, out), daemon=True).start()

    def add_questions(self, login, password, source, target, nums, points, subtype):
        self._run(shkolkovo.add_questions, login, password, source, target, nums, points, subtype)

    def upload_lessons(self, login, password, tasks):
        self._run(shkolkovo.upload_lessons, login, password, tasks)

    def _run(self, fn, *args):
        def job():
            log = egespace.Log(lambda msg: self._js("say", msg))
            try:
                fn(*args, log=log)
            except Exception as e:
                log.fail(str(e) if isinstance(e, (egespace.ApiError, ValueError)) else f"{type(e).__name__}: {e}")
            finally:
                self._js("done")
        threading.Thread(target=job, daemon=True).start()

    def backup_dir(self):
        return egespace.default_out()

    def restore(self, login, password, folder):
        threading.Thread(target=self._egespace, args=("restore", login, password, folder), daemon=True).start()

    def _egespace(self, what, login, password, arg, folder=""):
        say = lambda msg: self._js("say", msg)
        try:
            if what == "backup":
                ids = [int(x) for x in re.split(r"[^0-9]+", arg) if x]
                out = egespace.backup(login, password, ids, folder or None, say=say)
                try:
                    open_file(out)
                except OSError:
                    pass
            else:
                egespace.restore(login, password, arg, say=say)
        except Exception as e:
            say(f"Ошибка: {e if isinstance(e, (egespace.ApiError, ValueError)) else f'{type(e).__name__}: {e}'}")
        finally:
            self._js("done")

    def start(self, xlsx, pdf):
        threading.Thread(target=self._work, args=(xlsx, pdf), daemon=True).start()

    def _js(self, fn, *args):
        self._win.evaluate_js(f"{fn}({', '.join(json.dumps(a, ensure_ascii=False) for a in args)})")

    def _progress(self, i, n, msg):
        self._js("progress", i, n)
        self._js("say", f"[{i}/{n}] {msg}" if i else msg)

    _update = None

    def check_update(self):
        try:
            self._update = check_update()
        except Exception:  # нет интернета или GitHub недоступен — работаем на текущей версии
            self._update = None
        return {"build": self._update["build"], "auto": sys.platform == "win32"} if self._update else None

    def report(self, xlsx, first, last):
        threading.Thread(target=self._report, args=(xlsx, first, last), daemon=True).start()

    def _report(self, xlsx, first, last):
        try:
            self._js("say", "Читаю таблицу…")
            say = lambda msg: self._js("say", msg)
            out = make_report(xlsx, dt.date.fromisoformat(first), dt.date.fromisoformat(last), say=say)
            self._js("say", f"Отчет: {out}")
            try:
                open_file(out)
            except OSError as e:
                self._js("say", f"Не удалось открыть отчет: {e}")
        except Exception as e:
            self._js("say", f"Ошибка: {e if isinstance(e, ValueError) else f'{type(e).__name__}: {e}'}")
        finally:
            self._js("done")

    def update(self):
        threading.Thread(target=self._do_update, daemon=True).start()

    def _do_update(self):
        u = self._update
        if sys.platform != "win32":  # на macOS — страница релиза, .app заменить вручную
            return webbrowser.open(u["page"])
        exe = Path(sys.executable)
        new = exe.with_name(exe.name + ".new")
        try:
            self._js("say", f"Скачиваю сборку {u['build']}…")
            download(u["asset"], new, lambda done, total: self._js("progress", done, total))
            swap_after_exit(new, exe)
            self._js("say", "Готово — программа перезапустится уже новой версией.")
            os._exit(0)
        except Exception as e:
            self._js("say", f"Не удалось обновить: {e}")

    def _work(self, xlsx, pdf):
        web = None
        try:
            web = login_window(pdf)
            out, rows = run(xlsx, pdf, web, progress=self._progress)
            self._js("say", f"\n{summary(rows)}\nОтчёт: {out}")
            try:
                open_file(out)
            except OSError as e:
                self._js("say", f"Не удалось открыть отчёт: {e}")
        except Exception as e:
            self._js("say", f"\nОшибка: {e if isinstance(e, (LoginError, ValueError)) else f'{type(e).__name__}: {e}'}")
        finally:
            if web:
                web.close()
            self._js("done")


WEBVIEW2_URL = "https://developer.microsoft.com/microsoft-edge/webview2/"


def webview2_missing():
    """На Windows без WebView2 pywebview молча берёт движок IE, где окна не работают — лучше сказать сразу."""
    if sys.platform != "win32":
        return False
    from webview.platforms import winforms  # тот же выбор движка, что сделает webview.start()
    return winforms.renderer != "edgechromium"


def gui():
    if webview2_missing():
        import ctypes
        import webbrowser
        ctypes.windll.user32.MessageBoxW(
            0, "Для окна сайта нужен Microsoft Edge WebView2 Runtime. Сейчас откроется страница загрузки — "
               "установите «Evergreen Bootstrapper» и запустите программу снова.", "Сверка", 0x30)
        return webbrowser.open(WEBVIEW2_URL)
    api = Api()
    title = f"Сверка ExpertOnline — сборка {BUILD}" if BUILD else "Сверка ExpertOnline"
    api._win = webview.create_window(title, html=APP_HTML, js_api=api, width=820, height=580)
    webview.start(private_mode=True)


def main():
    if len(sys.argv) == 1:
        return gui()
    if sys.argv[1] == "--version-file":  # номер сборки в файл — им CI проверяет перезапуск после обновления
        return Path(sys.argv[2]).write_text(str(BUILD))
    if sys.argv[1] == "--swap-test":  # CI: собранный .exe подменяет себя и перезапускается, как при «Обновить»
        swap_after_exit(sys.argv[2], sys.executable, args=f'--version-file "{sys.argv[3]}"')
        os._exit(0)
    ap = argparse.ArgumentParser(description="Сверка «Получено от плательщиков» с xlsx и ежедневный отчет")
    ap.add_argument("--xlsx", required=True)
    ap.add_argument("--pdf", help="PDF со ссылками на группы — для сверки")
    ap.add_argument("--report-from", type=dt.date.fromisoformat, metavar="ГГГГ-ММ-ДД", help="отчет за день или с этого дня")
    ap.add_argument("--report-to", type=dt.date.fromisoformat, metavar="ГГГГ-ММ-ДД", help="…по этот день")
    ap.add_argument("--out")
    ap.add_argument("--site", default=SITE, help=argparse.SUPPRESS)  # для проверки на локальной копии
    a = ap.parse_args()
    for stream in (sys.stdout, sys.stderr):  # собранный .exe пишет в файл/pipe в cp1252 и падает на кириллице
        if stream is not None:
            stream.reconfigure(encoding="utf-8", errors="replace")
    if a.report_from:
        return print("Отчет:", make_report(a.xlsx, a.report_from, a.report_to, a.out))
    if not a.pdf:
        ap.error("нужен --pdf (сверка) или --report-from (отчет)")
    if webview2_missing():
        sys.exit(f"Нужен Microsoft Edge WebView2 Runtime: {WEBVIEW2_URL}")
    web, finished = login_window(a.pdf, a.site), threading.Event()

    def on_close():  # окно закрыли, Ctrl+C или Cmd+Q посреди сверки — не висим и не выходим «успешно»
        if not finished.is_set():
            print("Ошибка:", CLOSED, file=sys.stderr, flush=True)
            os._exit(1)

    web.win.events.closing += on_close
    web.win.events.closed += on_close

    def job():
        code = 1
        try:
            out, rows = run(a.xlsx, a.pdf, web, a.out, site=a.site,
                            progress=lambda i, n, msg: print(f"[{i}/{n}] {msg}" if i else msg, flush=True))
            print(summary(rows), "\nОтчёт:", out, flush=True)
            code = 0
        except Exception as e:
            print("Ошибка:", e, file=sys.stderr, flush=True)
        finally:
            finished.set()
            web.close()
            # ponytail: выходим сразу — на Windows destroy() окна, которое ещё грузится, теряется и start() не вернётся
            os._exit(code)

    webview.start(job, private_mode=True)


if __name__ == "__main__":
    main()
