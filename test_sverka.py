"""Самопроверка без настоящего сайта: синтетические xlsx/PDF + локальный сервер, похожий на HolliHop.

    python test_sverka.py                  # ядро сверки и отчёт
    python test_sverka.py --frozen <exe>   # собранное приложение целиком: окно сайта, вход, чтение страниц
    python test_sverka.py --swap           # Windows: подмена запущенного .exe при обновлении
    python test_sverka.py --swap-frozen <Sverka.exe>  # Windows: собранный .exe подменяет себя и реально перезапускается
"""
import difflib
import hashlib
import html
import os
import shutil
import subprocess
import sys
import tempfile
import threading
import time
import urllib.parse
import urllib.request
from datetime import datetime
from decimal import Decimal
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path

import docx
import openpyxl
from pypdf import PdfWriter
from pypdf.annotations import Link

import daily
import sverka

MONTHS = ["Сентябрь", "Октябрь"]
PAGES = {  # id группы → (название, «Получено от плательщиков»)
    1: ("26/27 сентябрь ГК ОБЩ БАЗА", "233&#160;596,00 руб."),       # совпадает
    2: ("26/27 октябрь ГК ОБЩ ПРОФИ 2", "75&#160;000,00 руб."),       # расхождение: в xlsx 75 105
    3: ("26/27 СВУ МВД ОБЩ", "275&#160;930,00 руб."),                 # сумма по месяцам
    4: ("26/27 октябрь ГК МАТ ТВИКС", "0,00"),                        # нет в xlsx, на сайте 0
    5: ("26/27 сентябрь ГК МАТ ТВИКС", "1&#160;000,00 руб."),         # нет в xlsx, на сайте деньги
    6: ("25/26 на пересдачу: математика / базовый минимум", "0,00"),   # две группы → одна ячейка
    7: ("25/26 на пересдачу: математика / роскошный максимум", "43&#160;330,00 руб."),
    8: ("С-П", "1&#160;500,00 руб."),                                  # алиас → «Стартер-пак»
    9: ("26/27 сентябрь ГК ОБЩ ПРО+ПА", "98&#160;671,00 руб."),
}
BROKEN = 10  # страница без «Получено от плательщиков»
PAYMENTS = {  # id группы → строки списка «Оплаты»: плательщик (с инициалами, как на сайте), сумма, статус
    2: [("Иванова А. Б.", "30&#160;000,00 руб.", "Оплачен"), ("Петров В. Г.", "45&#160;000,00 руб.", ""),
        ("Сидорова Е. Е.", "999,00 руб.", "Счёт: №1 Не оплачен")],  # неоплаченный счёт не считается
    5: [("Смирнов П. Р.", "1&#160;000,00 руб.", "Оплачен")],
    8: [("Евсеенко К. В.", "1&#160;000,00 руб.", "Оплачен"), ("Кузнецова В. А.", "500,00 руб.", "Оплачен")],
    9: [("Орлова О. О.", "48&#160;671,00 руб.", "Оплачен"),  # сумма группы сходится, а по ученикам — нет
        ("Белов Б. Б. заказчик: Белова Анна Ивановна", "50&#160;000,00 руб.", "Оплачен")],
}
PAGE_SIZE = {2: 1}  # по одной оплате на страницу — проверка листания списка
LOGIN, PASSWORD = "test", "test-pass"
FOREIGN = "http://evil.example/Learner/Group/1"  # первая ссылка PDF на чужой хост — пароль туда уходить не должен


def make_xlsx(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "ОБЩ ГК "  # 26/27, с пробелом в конце, как в настоящем файле
    ws.append(["Дата", "ФИО", *MONTHS, *MONTHS, "Тариф"])
    summary = [("ГК база", 233596, 174291), ("ГК профи 2", 0, 75105),
               ("ГК про+ПА", 98671, 86181), ("СВУ МВД", 275930, 0)]
    for label, sep, oct_ in summary:
        ws.append(["ФАКТ", None, 1, 2, label, label])
        ws.append([None, None, "оплат", "оплат", sep, oct_])  # «оплат» над числом не должно стать подписью
    ws["F5"] = "=SUMIF(F11:F30,F4,D11:D30)"  # как в настоящем файле: по формуле видно, каких учеников она складывает
    for r, (name, paid) in enumerate([("Иванова Анна Борисовна", 30000), ("Петров Василий Глебович", 45105)], start=11):
        ws.cell(r, 2, name)
        ws.cell(r, 4, paid)
        ws.cell(r, 6, "ГК профи 2")
    ws["E7"] = "=SUMIF(E11:E30,E6,C11:C30)"  # про+ПА сентябрь: 50000 + 48671 = 98671, как на сайте
    for r, (name, paid) in enumerate([("Орлова Ольга Олеговна", 50000), ("Белов Борис Борисович", 48671)], start=13):
        ws.cell(r, 2, name)
        ws.cell(r, 3, paid)
        ws.cell(r, 5, "ГК про+ПА")
    old = wb.create_sheet("ОБЩ ГК")  # прошлый год, другие тарифы — не должен мешать
    old.append(["Дата", "ФИО", *MONTHS, *MONTHS])
    old.append([None, None, None, None, "ГК pro+ПА", "ГК pro+ПА"])
    old.append([None, None, None, None, 1, 1])
    mat = wb.create_sheet("МАТ ГК")
    mat.append(["Дата", "ФИО", *MONTHS, *MONTHS])
    mat.append([None, None, None, None, "ГК база", "ГК база"])
    mat.append([None, None, None, None, 0, 12502])
    for name, rows in [("МАТ спецкурсы", [("Все", 43330), ("На пересдачу", 43330)]),
                       ("ОБЩ спецкурсы", [("Все", 1000), ("Стартер-пак", "=SUMIF(D5:D10,D3,C5:C10)")]),
                       ("ИСТ спецкурсы", [("Все", 500), ("для него в тарифе пустым оставлять", "Стартер-пак")])]:
        s = wb.create_sheet(name)
        s.append(["Дата", "ФИО", "Сумма", "Курс"])
        for course, total in rows:
            s.append(["Количество", 1, total, course])  # у ИСТ сумма съехала — как в настоящем файле
    for r, name in enumerate(["Евсеенко Кристина Вячеславовна", "Кузнецова Вера Андреевна"], start=5):
        for c, v in zip((2, 3, 4), (name, 500, "Стартер-пак")):
            wb["ОБЩ спецкурсы"].cell(r, c, v)
    for c, v in zip((2, 3, 4), ("Евсеенко Кристина Вячеславовна", 500, "Стартер-пак")):  # та же ученица на листе ИСТ
        wb["ИСТ спецкурсы"].cell(4, c, v)
    wb.create_sheet("ЗП 26-27").append(["факт", 1, 2, "ГК база"])
    wb.save(path)


def payments_page(gid, page):
    rows, size = PAYMENTS.get(gid, []), PAGE_SIZE.get(gid, 15)
    pages = max(1, -(-len(rows) // size))
    trs = "".join(f'<tr><td>01.10.2026</td><td><a href="/Payment/Payer/1">{html.escape(p)}</a></td><td>{a}</td>'
                  f'<td>{html.escape(st)}</td><td><a href="/Payment/RemoveLearnerPayment/{gid}">удалить</a></td></tr>'
                  for p, a, st in rows[page * size:(page + 1) * size])
    pager = "".join(f'<a class="page-link" href="/Payment/SearchLearnerPayments?LearnerId={gid}&Page={i}">{i + 1}</a>'
                    for i in range(pages))
    return ("<title>Оплаты</title><table><tr><th>Дата добавления</th><th>Плательщик</th><th>Размер оплаты</th>"
            f"<th>Статус</th><th></th></tr>{trs}</table><div>{pager}</div><p>Всего найдено {len(rows)} оплат</p>")


def static_payments(url):  # то же, что отдаёт фейковый сайт, — для проверки без окна
    rows = [["01.10.2026", p, html.unescape(a), st] for p, a, st in PAYMENTS.get(int(url.rstrip("/").split("/")[-1]), [])]
    return sverka.payment_rows(["Дата добавления", "Плательщик", "Размер оплаты", "Статус"], rows), "Всего найдено"


def make_pdf(path, base):
    w = PdfWriter()
    w.add_blank_page(600, 800)
    urls = [f"{base}/Learner/Group/{i}" if i != 8 else f"{base}/Learner/8"  # короткая: сайт редиректит на Group/8
            for i in [*PAGES, BROKEN]]
    urls.insert(1, urls[0])  # повтор ссылки — должна посчитаться один раз
    urls.append(f"{base}/Learner/Group/8")  # та же группа 8 в другом написании
    urls.insert(0, FOREIGN)
    urls.append("https://example.com/not-a-group")
    for i, url in enumerate(urls):
        w.add_annotation(0, Link(rect=(10, 10 + 15 * i, 200, 22 + 15 * i), url=url))
    w.write(path)


# Тестовый сайт входит сам вместо человека: форма с тестовым логином отправляется скриптом страницы.
LOGIN_PAGE = ('<form method="post" action="{action}"><input name="LogLogin" value="{login}">'
              '<input name="LogPassword" type="password" value="{password}"></form>'
              "<script>document.forms[0].submit()</script>")


class FakeHolliHop(BaseHTTPRequestHandler):
    log = []

    def log_message(self, *a):
        pass

    def _send(self, code, body="", headers=()):
        self.send_response(code)
        for k, v in headers:
            self.send_header(k, v)
        data = body.encode()
        self.send_header("Content-Type", "text/html; charset=utf-8")
        self.send_header("Content-Length", str(len(data)))
        self.end_headers()
        self.wfile.write(data)

    def do_POST(self):
        self.log.append(("POST", self.path))
        u = urllib.parse.urlsplit(self.path)
        form = urllib.parse.parse_qs(self.rfile.read(int(self.headers["Content-Length"])).decode())
        if u.path == "/Account/LogOn" and form.get("LogLogin") == [LOGIN] and form.get("LogPassword") == [PASSWORD]:
            back = urllib.parse.parse_qs(u.query).get("ReturnUrl", ["/"])[0]
            return self._send(302, headers=[("Location", back), ("Set-Cookie", ".ASPXAUTH=ok; path=/; HttpOnly")])
        self._send(200, "<p>Неверный логин или пароль</p>")

    def do_GET(self):
        self.log.append(("GET", self.path))
        path = urllib.parse.urlsplit(self.path).path
        if path == "/Account/LogOn":
            return self._send(200, LOGIN_PAGE.format(action=html.escape(self.path), login=LOGIN, password=PASSWORD))
        if path in ("/", "/favicon.ico"):
            return self._send(200, "<p>Главная</p>")
        if path == "/asset":  # «файл релиза» для проверки обновления
            return self._send(200, "new")
        if ".ASPXAUTH=ok" not in (self.headers.get("Cookie") or ""):
            return self._send(302, headers=[("Location", "/Account/LogOn?ReturnUrl=" + urllib.parse.quote(path))])
        if path == "/Payment/SearchLearnerPayments":
            q = urllib.parse.parse_qs(urllib.parse.urlsplit(self.path).query)
            return self._send(200, payments_page(int(q["LearnerId"][0]), int(q.get("Page", ["0"])[0])))
        parts = path.strip("/").split("/")
        if parts[:1] == ["Learner"] and parts[-1].isdigit() and len(parts) == 2:
            return self._send(302, headers=[("Location", f"/Learner/Group/{parts[-1]}")])
        gid = int(parts[-1])
        if gid == BROKEN:
            return self._send(200, "<title>Группа без статистики</title>")
        name, received = PAGES[gid]
        self._send(200, f"<title>Группа {html.escape(name)}</title><header><x-caption class=\"col-sm\">\r\n\r\n"
                        f"{html.escape(name)}                    <x-status>в работе</x-status></x-caption></header>"
                        f"<table><tr class=\"TrNoTopBorders\">\r\n<td>Получено от плательщиков:</td>\r\n"
                        f"<td><a href=\"/Payment/SearchLearnerPayments?LearnerId={gid}\">{received}</a></td></tr></table>")


def setup():
    server = ThreadingHTTPServer(("127.0.0.1", 0), FakeHolliHop)
    threading.Thread(target=server.serve_forever, daemon=True).start()
    base = f"http://127.0.0.1:{server.server_port}"
    tmp = Path(tempfile.mkdtemp())
    xlsx, pdf, out = tmp / "t.xlsx", tmp / "t.pdf", tmp / "r.docx"
    make_xlsx(xlsx)
    make_pdf(pdf, base)
    FakeHolliHop.log.clear()
    return server, base, xlsx, pdf, out


def check_report(out):
    doc = docx.Document(out)
    text = "\n".join(c.text for t in doc.tables for row in t.rows for c in row.cells)
    assert "РАСХОЖДЕНИЕ" in text and "75\u00a0000,00" in text and "-105,00" in text, text[:500]
    paras = "\n".join(p.text for p in doc.paragraphs)
    assert "Расхождения по ученикам" in paras and "Петров Василий Глебович" in text and "45\u00a0105,00" in text, text[-800:]
    assert "Ученики с расхождением" in text and "Петров В. Г.: сайт 45\u00a0000,00 / xlsx 45\u00a0105,00" in text
    assert "РАСХОЖДЕНИЕ по ученикам" in text and "Белов Б. Б. (заказчик: Белова Анна Ивановна)" in text
    assert "Евсеенко К. В." in text and "Смирнов П. Р." in text and "Сидорова" not in text, "неоплаченный счёт не в отчёте"
    assert FOREIGN in paras and "«ИСТ спецкурсы», строка 3: в C вместо суммы «Стартер-пак»" in paras, paras


def check_read_only():
    """На сайте только GET страниц групп и входа; POST — только форма входа (её отправляет страница сайта)."""
    log = FakeHolliHop.log
    assert log, "к сайту не обращались"
    assert {m for m, _ in log} <= {"GET", "POST"}, log
    assert all(urllib.parse.urlsplit(p).path == "/Account/LogOn" for m, p in log if m == "POST"), log
    assert all(urllib.parse.urlsplit(p).path in ("/", "/favicon.ico", "/Account/LogOn", "/Payment/SearchLearnerPayments")
               or p.startswith("/Learner/") for m, p in log if m == "GET"), log
    assert not any("Remove" in p or "Edit" in p for _, p in log), "кнопки «удалить»/«изменить» на странице оплат не нажимаются"


def main():
    server, base, xlsx, pdf, out = setup()
    opener = urllib.request.build_opener(urllib.request.HTTPCookieProcessor())  # тестовый клиент вместо окна сайта
    opener.open(base + "/Account/LogOn", urllib.parse.urlencode({"LogLogin": LOGIN, "LogPassword": PASSWORD}).encode())
    fetch = lambda url: opener.open(base + sverka.group_path(url, base)).read().decode()

    rows, links, skipped, problems = sverka.reconcile(xlsx, pdf, fetch, site=base, payments=static_payments)
    sverka.write_report(out, rows, links, skipped, problems, xlsx, pdf)
    by_name = {g.name or g.url: r for r in rows for g in r.groups}
    status = {name: r.status for name, r in by_name.items()}
    assert status["26/27 сентябрь ГК ОБЩ БАЗА"] == "ok"
    pa = by_name["26/27 сентябрь ГК ОБЩ ПРО+ПА"]
    assert pa.cells == ["«ОБЩ ГК »!E7"], "старый лист «ОБЩ ГК» не должен мешать"
    assert (pa.status, pa.diff) == ("students", Decimal("0.00")), "сумма сходится — но два ученика гасят друг друга"
    assert pa.students == [("Белов Б. Б. (заказчик: Белова Анна Ивановна)", "Белов Борис Борисович", Decimal(50000), Decimal("48671.00"), ""),
                           ("Орлова О. О.", "Орлова Ольга Олеговна", Decimal(48671), Decimal("50000.00"), "")], pa.students
    assert status["26/27 сентябрь ГК ОБЩ БАЗА"] == "ok" and "не формулой SUMIF" in by_name["26/27 сентябрь ГК ОБЩ БАЗА"].students_note
    r = by_name["26/27 октябрь ГК ОБЩ ПРОФИ 2"]
    assert (r.status, r.site, r.xlsx, r.diff, r.cells) == ("diff", Decimal("75000"), Decimal("75105.00"), Decimal("-105.00"), ["«ОБЩ ГК »!F5"])
    assert by_name["26/27 СВУ МВД ОБЩ"].xlsx == Decimal("275930") and status["26/27 СВУ МВД ОБЩ"] == "ok"
    assert status["26/27 октябрь ГК МАТ ТВИКС"] == "empty"
    assert status["26/27 сентябрь ГК МАТ ТВИКС"] == "missing"
    retake = by_name["25/26 на пересдачу: математика / роскошный максимум"]
    assert retake is by_name["25/26 на пересдачу: математика / базовый минимум"] and retake.status == "ok"
    sp = by_name["С-П"]
    assert (sp.status, sp.diff, sp.cells) == ("diff", Decimal("500.00"), ["«ОБЩ спецкурсы»!C3"])
    # кто из учеников дает расхождение: инициалы на сайте ↔ полное ФИО в xlsx
    assert sp.students == [("Евсеенко К. В.", "Евсеенко Кристина Вячеславовна", Decimal("1000"), Decimal("500.00"),
                            "«ИСТ спецкурсы» стр. 4")], sp.students  # вторая половина оплаты — на другом листе
    assert r.students == [("Петров В. Г.", "Петров Василий Глебович", Decimal("45000"), Decimal("45105.00"), "")]
    assert by_name["26/27 сентябрь ГК МАТ ТВИКС"].students == [("Смирнов П. Р.", "—", Decimal("1000"), Decimal("0"), "")]
    assert status[f"{base}/Learner/Group/{BROKEN}"] == "error"
    assert sum(len(r.groups) for r in rows) == len(PAGES) + 1  # + битая; повтор ссылки — один раз
    check_read_only()
    for bad in (base + "/Payment/Delete/1", FOREIGN, "https://expertonline.t8s.ru/Learner/Group/1"):
        try:
            sverka.group_path(bad, base)
            raise AssertionError(f"group_path должен отказать: {bad}")
        except ValueError:
            pass
    check_report(out)
    assert sverka.parse_money("−1\u00a0234,50 руб.") == Decimal("-1234.50")
    assert sverka.parse_group_name("26/27 октябрьГК ОБЩ ПРО + ВГ") == ("ОБЩ", "Октябрь", "про+вг")
    check_updates(base)
    check_daily()
    check_names()
    server.shutdown()
    print("ok:", sverka.summary(rows))


def make_daily_xlsx(path):
    """Мини-таблица со всеми правилами отчета: возврат выше диапазона учеников, СВУ МВД со «штампом» даты
    после дня отчета, латинская O в «ПРO+ПА», ученик после дня отчета, стартер-пак нарастающим итогом."""
    d = lambda day, month=10: datetime(2026, month, day)
    wb = openpyxl.Workbook()
    zp = wb.active
    zp.title = "ЗП 26-27"
    zp["A3"], zp["A6"] = d(1), "Итого"
    for col, (title, plan) in zip((2, 5, 8, 11), [("обществознание", 500000), ("история", 100000),
                                                  ("русский язык", 50000), ("математика", 50000)]):
        zp.cell(3, col, title)
        zp.cell(6, col + 2, plan)
    log = {"Обществознание": [(d(1), 5290), (d(1), 2145), (d(2), 4790), (d(30, 9), 999)],
           "История": [(d(2), 4496)], "Русский язык": [], "Математика": []}
    for col, (title, rows) in zip((20, 25, 30, 35), log.items()):
        zp.cell(1, col, title)
        for i, h in enumerate(("Дата", "ФИО", "Сумма")):
            zp.cell(2, col + i, h)
        for r, (day, amount) in enumerate(rows, start=3):
            zp.cell(r, col, day)
            zp.cell(r, col + 1, "Ученик")
            zp.cell(r, col + 2, amount)
    for subj in daily.SUBJ:
        ws = wb.create_sheet(f"{subj} ГК")
        ws.append(["Дата", "ФИО", "Сентябрь", "Октябрь", "Сентябрь", "Октябрь", "Тариф"])
        summary = [("Сумма", "=SUM(C8:D100)"), ("Количество", "=COUNTA(B10:B100)"), ("БАЗА", "=COUNTIF(G10:G100,G4)"),
                   ("ПРO+ПА", "=COUNTIF(G10:G100,G5)"), ("СВУ МВД", "=COUNTIF(G9:G100,G6)")]
        for r, (label, formula) in enumerate(summary, start=2):
            ws.cell(r, 7, label)
            ws.cell(r, 8, formula)
        if subj == "ОБЩ":
            for r, row in {8: (d(1, 9), "Возврат", 1000, None, "БАЗА"), 9: (d(8), "Курсант", 2000, None, "СВУ МВД"),
                           10: (d(15, 9), "Ученик 1", 3000, 3000, "БАЗА"), 11: (d(1), "Ученик 2", None, 5290, "ПРО+ПА"),
                           12: (d(2), "Ученик 3", None, 4790, "БАЗА")}.items():
                for c, v in zip((1, 2, 3, 4, 7), row):
                    ws.cell(r, c, v)
    sp = wb.create_sheet("ОБЩ спецкурсы")
    sp.append(["Дата", "ФИО", "Сумма", "Курс"])
    sp.append(["Количество", 1, 500, "Стартер-пак"])
    sp.append([d(10, 9), "А", 500, "Стартер-пак"])
    sp.append([d(5), "Б", 500, "Стартер-пак"])
    wb.save(path)


DAILY_OCT1 = """📌 Отчет за 1 октября:

Всего обработано: 

🔵 ОБЩЕСТВОЗНАНИЕ

Получено оплат: 2
Сумма оплат за сегодня: 7.435
Процент выполнения плана: 7.435 от 500.000 - 1.49%

Годовой курс: 2
В деньгах: 14.290
БАЗА 1 / ПРО+ПА 1 / СВУ МВД 1

Стартер-пак: 1

🟣 ИСТОРИЯ

Получено оплат: 0
Сумма оплат за сегодня: 0
Процент выполнения плана: 0.00 от 100.000 - 0.00%

Годовой курс: 0
В деньгах: 0
БАЗА 0 / ПРО+ПА 0 / СВУ МВД 0

🟠 РУССКИЙ ЯЗЫК

Получено оплат: 0
Сумма оплат за сегодня: 0
Процент выполнения плана: 0.00 от 50.000 - 0.00%

Годовой курс: 0
В деньгах: 0
БАЗА 0 / ПРО+ПА 0 / СВУ МВД 0

🔵 МАТЕМАТИКА

Получено оплат: 0
Сумма оплат за сегодня: 0
Процент выполнения плана: 0.00 от 50.000 - 0.00%

Годовой курс: 0
В деньгах: 0
БАЗА 0 / ПРО+ПА 0 / СВУ МВД 0"""


def check_daily():
    path = Path(tempfile.mkdtemp()) / "Ученики.xlsx"
    make_daily_xlsx(path)
    book = daily.Book(path)
    got = book.report(datetime(2026, 10, 1).date())
    assert got == DAILY_OCT1, "\n".join(difflib.unified_diff(DAILY_OCT1.splitlines(), got.splitlines(), lineterm=""))
    oct2 = book.report(datetime(2026, 10, 2).date())  # день спустя: ученик 3 добавился, процент — с начала месяца
    assert "Получено оплат: 1\nСумма оплат за сегодня: 4.790\nПроцент выполнения плана: 12.225 от 500.000 - 2.45%" in oct2
    assert "Годовой курс: 3\nВ деньгах: 19.080\nБАЗА 2 / ПРО+ПА 1 / СВУ МВД 1" in oct2
    assert daily.days(datetime(2026, 10, 3).date(), datetime(2026, 10, 1).date())[0].day == 1  # перепутанные даты
    out = sverka.make_report(path, datetime(2026, 10, 1).date(), datetime(2026, 10, 2).date(), say=lambda m: None)
    assert [p.text for p in docx.Document(out).paragraphs][:1] == ["📌 Отчет за 1 октября:"]
    try:
        book.report(datetime(2026, 11, 1).date())
        raise AssertionError("месяц без плана должен давать понятную ошибку")
    except ValueError as e:
        assert "нет плана за 11.2026" in str(e)


def check_names():
    # в xlsx имя первым, на сайте «Фамилия И. О.» — это один человек, а не две гасящие друг друга строки
    row = sverka.Row([sverka.Group("u", "g", Decimal(1500))], "diff", hits=[sverka.Entry(
        "ОБЩ ГК ", "ОБЩ", "Сентябрь", "профи1", Decimal(1000), "L5",
        [("Сабрина Русу Олеговна", Decimal(500), 29), ("Гусев Никита Дмитриевич", Decimal(500), 30)])])
    people = {("ситникова", "в"): [("ОБЩ ГК", 537, "Ситникова Василиса Дмитриевна")],  # однофамилица, прошлый год
              ("гусев", "н"): [("ОБЩ ГК", 40, "Гусев Никита Дмитриевич"), ("РУС ГК", 12, "Гусев Никита Дмитриевич")]}
    row.hits[0].students.append(("Карамышеаа Софья Евгеньевна", Decimal(6190), 31))  # опечатка в xlsx
    pays = [("Русу С. О.", Decimal(500)), ("Гусев Н. Д.", Decimal(1500)), ("Ситникова В. М.", Decimal(100)),
            ("Карамышева С. Е.", Decimal(6190))]
    sverka.student_diff(row, lambda url: (pays, ""), people, skip_sheets={"ОБЩ ГК"})
    assert row.students == [("Гусев Н. Д.", "Гусев Никита Дмитриевич", Decimal(1500), Decimal(500), "«РУС ГК» стр. 12"),
                            ("Ситникова В. М.", "—", Decimal(100), Decimal(0), "")], row.students
    assert row.typos == [("Карамышева С. Е.", "Карамышеаа Софья Евгеньевна")], "опечатка — не денежное расхождение"
    assert sverka.same_person("Евсеенко К. В.", "Евсеенко Кристина Вячеславовна")
    assert not sverka.same_person("Ситникова Василиса Михайловна", "Ситникова Василиса Дмитриевна")
    assert sverka.name_key("Cидорова Анна") == sverka.name_key("Сидорова А.")  # латинская C в xlsx


def check_updates(base):
    asset = {"name": sverka.asset_name(), "size": 3, "browser_download_url": base + "/asset",
             "digest": "sha256:" + hashlib.sha256(b"new").hexdigest()}
    release = {"tag_name": "build-7", "html_url": "https://example.com/r/7", "assets": [asset]}
    assert sverka.find_update(release, 6)["build"] == 7
    assert sverka.find_update(release, 7) is None, "та же сборка — не обновление"
    assert sverka.find_update(dict(release, tag_name="v1.0"), 1) is None
    assert sverka.find_update(dict(release, assets=[]), 6) is None, "нет файла для этой ОС"
    assert sverka.check_update(build=0) is None, "из исходников обновления не проверяются (и сеть не нужна)"
    dest = Path(tempfile.mkdtemp()) / "Sverka.exe.new"
    sverka.download(asset, dest)
    assert dest.read_bytes() == b"new"
    for bad in (dict(asset, digest="sha256:" + "0" * 64), dict(asset, size=4)):
        try:
            sverka.download(bad, dest)
            raise AssertionError("битый файл должен отбрасываться")
        except ValueError:
            assert not dest.exists()


def swap():
    """Windows: подмена запущенного .exe скриптом обновления. Вместо Sverka — копии системных ping.exe и cmd.exe."""
    d, sysdir = Path(tempfile.mkdtemp()), Path(os.environ["SystemRoot"]) / "System32"
    exe, new = d / "Sverka.exe", d / "Sverka.exe.new"
    shutil.copy(sysdir / "PING.EXE", exe)
    shutil.copy(sysdir / "cmd.exe", new)
    old = subprocess.Popen([str(exe), "-n", "4", "127.0.0.1"], stdout=subprocess.DEVNULL)  # «старая версия» ещё работает
    sverka.swap_after_exit(new, exe, args="/c mkdir restarted")
    old.wait()
    for _ in range(60):
        if (d / "restarted").exists():
            break
        time.sleep(1)
    assert (d / "restarted").exists() and not new.exists(), list(d.iterdir())
    assert exe.read_bytes() == (sysdir / "cmd.exe").read_bytes()
    print("ok swap:", d)


def swap_frozen(built):
    """Собранный onefile .exe подменяет себя копией и перезапускается — ловит «Failed to load Python DLL»."""
    d = Path(tempfile.mkdtemp())
    exe, new, marker = d / "Sverka.exe", d / "Sverka.exe.new", d / "build.txt"
    shutil.copy(built, exe)
    shutil.copy(built, new)
    subprocess.run([str(exe), "--swap-test", str(new), str(marker)], timeout=120)
    for _ in range(90):
        if marker.exists():
            break
        time.sleep(1)
    assert marker.exists() and not new.exists(), ("новая версия не запустилась после подмены", list(d.iterdir()))
    print("ok swap-frozen: перезапустилась сборка", marker.read_text())


def frozen(exe):
    """Собранное приложение в консольном режиме: открывает окно сайта, тестовый сайт входит сам, дальше сверка."""
    server, base, xlsx, pdf, out = setup()
    logs = out.with_suffix(".log")
    with open(logs, "w") as f:  # не pipe: onefile-.exe держит его дочерним процессом, и ожидание вечное
        p = subprocess.Popen([exe, "--xlsx", str(xlsx), "--pdf", str(pdf), "--out", str(out), "--site", base],
                             stdout=f, stderr=subprocess.STDOUT)
        try:
            code = p.wait(timeout=180)
        except subprocess.TimeoutExpired:
            if sys.platform == "win32":
                subprocess.run(["taskkill", "/F", "/T", "/PID", str(p.pid)], check=False)
            p.kill()
            code = "timeout"
    assert code == 0 and out.exists(), (code, logs.read_text(errors="replace")[-3000:], FakeHolliHop.log)
    check_report(out)
    check_read_only()
    server.shutdown()
    daily_xlsx, daily_out = out.with_name("Ученики.xlsx"), out.with_name("Отчет.docx")  # отчет внутри сборки
    make_daily_xlsx(daily_xlsx)
    subprocess.run([exe, "--xlsx", str(daily_xlsx), "--report-from", "2026-10-01", "--out", str(daily_out)], timeout=120)
    assert "\n".join(p.text for p in docx.Document(daily_out).paragraphs) == DAILY_OCT1
    print("ok frozen:", exe, "| запросов к сайту:", len(FakeHolliHop.log), "| отчет совпал")


if __name__ == "__main__":
    if sys.argv[1:2] == ["--frozen"]:
        frozen(sys.argv[2])
    elif sys.argv[1:2] == ["--swap"]:
        swap()
    elif sys.argv[1:2] == ["--swap-frozen"]:
        swap_frozen(sys.argv[2])
    else:
        main()
