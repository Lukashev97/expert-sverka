"""Ежедневный отчёт руководителю из таблицы учеников — тот же текст, что пишется руками.

Откуда цифры (подобрано так, чтобы отчёт за 1 октября совпал с ручным байт в байт):
- «Получено оплат», «Сумма оплат за сегодня» — журнал продаж справа на листе «ЗП 26-27», строки с этой датой.
- «Процент выполнения плана» — продажи с начала месяца по эту дату от плана месяца (блоки менеджеров слева, строка «Итого»).
- «Годовой курс», тарифы, «В деньгах» — сводка наверху листа «<ПРЕДМЕТ> ГК»: её же формулы COUNTA / COUNTIF / SUM,
  только по строкам с Датой не позже дня отчёта. Красные/серые (возвраты, ушедшие) стоят выше этих диапазонов.
  СВУ МВД — отдельная когорта, по дате не фильтруется.
- «Стартер-пак» — нарастающим итогом с листа «<ПРЕДМЕТ> спецкурсы»; строки нет, если предмет его не продаёт.
"""
import datetime as dt
import re
import warnings
from collections import Counter
from decimal import ROUND_HALF_UP, Decimal

import docx
import openpyxl
from docx.enum.text import WD_BREAK
from openpyxl.utils import column_index_from_string, get_column_letter

SUBJ = ["ОБЩ", "ИСТ", "РУС", "МАТ"]
LOG_TITLE = {"ОБЩ": "Обществознание", "ИСТ": "История", "РУС": "Русский язык", "МАТ": "Математика"}
HEAD = {"ОБЩ": "🔵 ОБЩЕСТВОЗНАНИЕ", "ИСТ": "🟣 ИСТОРИЯ", "РУС": "🟠 РУССКИЙ ЯЗЫК", "МАТ": "🔵 МАТЕМАТИКА"}
GEN = "января февраля марта апреля мая июня июля августа сентября октября ноября декабря".split()
SVU = "СВУ МВД"
ZP = "ЗП 26-27"
LAT2CYR = str.maketrans("ABCEHKMOPTXYaceopxy", "АВСЕНКМОРТХУасеорху")  # в листе «ПРO+ПА» с латинской O
REFUND_FILLS = {"FFFFCFC9", "FFCFCECE", "FFD0CECE", "FFB7B7B7"}
RNG = r"\$?([A-Z]+)\$?(\d+):\$?([A-Z]+)\$?(\d+)"


def as_date(v):
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    if isinstance(v, str):  # бывает вписано руками: «31 08.2026»
        m = re.fullmatch(r"\s*(\d{1,2})[\s.,/-]+(\d{1,2})[\s.,/-]+(\d{4})\s*", v)
        if m:
            try:
                return dt.date(int(m[3]), int(m[2]), int(m[1]))
            except ValueError:
                return None
    return None


def num(v):
    return v if isinstance(v, (int, float)) and not isinstance(v, bool) else 0


def norm(s):
    return str(s).strip().translate(LAT2CYR).casefold()


def money(x):
    return f"{round(x):,}".replace(",", ".")


def _fill(c):
    f = c.fill
    return f.fgColor.rgb if f is not None and f.fill_type and f.fgColor.type == "rgb" else None


class Book:
    def __init__(self, path):
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            self.v = openpyxl.load_workbook(path, data_only=True)  # значения
            self.f = openpyxl.load_workbook(path)  # формулы сводок — по ним берём диапазоны
        if ZP not in self.v.sheetnames:
            raise ValueError(f"в таблице нет листа «{ZP}»")
        self.zp = self.v[ZP]
        self.logs = {s: self._log(s) for s in SUBJ}
        self.plans = self._plans()
        self.last_log_day = max((d for rows in self.logs.values() for _, d, _ in rows), default=None)

    # ---------- «ЗП 26-27»: журнал продаж и планы ----------

    def _log(self, s):
        ws = self.zp
        c0 = next((c for c in range(1, ws.max_column + 1) if str(ws.cell(1, c).value or "").strip() == LOG_TITLE[s]), None)
        if c0 is None:
            raise ValueError(f"на листе «{ZP}» нет журнала «{LOG_TITLE[s]}»")
        hdr = {}
        for c in range(c0, c0 + 9):  # первое вхождение: соседний предмет может начинаться ближе 9 колонок
            hdr.setdefault(str(ws.cell(2, c).value or "").strip(), c)
        cd, cs = hdr["Дата"], hdr["Сумма"]
        out = []
        for r in range(3, ws.max_row + 1):
            d, a = as_date(ws.cell(r, cd).value), ws.cell(r, cs).value
            if d and isinstance(a, (int, float)):
                out.append((r, d, a))
        return out

    def _plans(self):
        """{(год, месяц): {предмет: план}}. Месяц блока — по датам журнала в его строках (в шапке бывают опечатки)."""
        ws = self.zp
        heads = [r for r in range(1, ws.max_row + 1) if isinstance(ws.cell(r, 1).value, dt.datetime)]
        months = {}
        for rows in self.logs.values():
            for r, d, _ in rows:
                months.setdefault(r, []).append((d.year, d.month))
        out = {}
        for i, h in enumerate(heads):
            end = heads[i + 1] if i + 1 < len(heads) else ws.max_row + 1
            seen = Counter(ym for r in range(h, end) for ym in months.get(r, []))
            ym = seen.most_common(1)[0][0] if seen else (ws.cell(h, 1).value.year, ws.cell(h, 1).value.month)
            itogo = next((r for r in range(h, end) if ws.cell(r, 1).value == "Итого"), None)
            if itogo is None:
                continue
            plan = {}
            for s in SUBJ:
                col = next((c for c in range(2, 40) if norm(ws.cell(h, c).value or "") == norm(LOG_TITLE[s])), None)
                plan[s] = num(ws.cell(itogo, col + 2).value) if col else 0
            out[ym] = plan
        return out

    # ---------- «<ПРЕДМЕТ> ГК» ----------

    def _gk(self, wb, s):
        ws = next((w for w in wb.worksheets if w.title.strip() == f"{s} ГК" and w.sheet_state == "visible"), None)
        if ws is None:
            raise ValueError(f"в таблице нет листа «{s} ГК»")
        return ws

    def gk(self, s, day):
        wv, wf = self._gk(self.v, s), self._gk(self.f, s)
        hdr = [c.value for c in wv[1]]
        if "Тариф" not in hdr:
            raise ValueError(f"на листе «{wv.title}» нет колонки «Тариф»")
        tc = hdr.index("Тариф") + 1
        letter = get_column_letter(tc)
        lab = {str(wf.cell(r, tc).value).strip(): r for r in range(2, 20) if wf.cell(r, tc).value}
        try:
            f_sum = wf.cell(lab["Сумма"], tc + 1).value.replace(" ", "")
            f_cnt = wf.cell(lab["Количество"], tc + 1).value.replace(" ", "")
        except (KeyError, AttributeError):
            raise ValueError(f"на листе «{wv.title}» не найдена сводка «Сумма» / «Количество»") from None

        def keep(r):
            if norm(wv.cell(r, tc).value or "") == norm(SVU):
                return True
            d = as_date(wv.cell(r, 1).value)
            return d is None or d <= day

        m = re.fullmatch(r"=SUM\(" + RNG + r"\)", f_sum)
        c0, r0, c1 = column_index_from_string(m[1]), int(m[2]), column_index_from_string(m[3])
        total = sum(num(wv.cell(r, c).value) for r in range(r0, wv.max_row + 1) if keep(r) for c in range(c0, c1 + 1))

        m = re.fullmatch(r"=COUNTA\(" + RNG + r"\)(?:-(\d+))?", f_cnt)  # у РУС в формуле вбито «-1» — повторяем
        students = sum(1 for r in range(int(m[2]), wv.max_row + 1)
                       if str(wv.cell(r, 2).value or "").strip() and keep(r)) - int(m[5] or 0)

        tariffs = []
        for r in range(lab["Количество"] + 1, lab["Количество"] + 12):
            mm = re.fullmatch(r"=COUNTIF\(" + RNG + r"," + letter + r"(\d+)\)", str(wf.cell(r, tc + 1).value or "").replace(" ", ""))
            if not mm:
                break
            crit = wv.cell(r, tc).value
            k = sum(1 for rr in range(int(mm[2]), wv.max_row + 1) if norm(wv.cell(rr, tc).value or "") == norm(crit) and keep(rr))
            tariffs.append((str(crit).strip().translate(LAT2CYR), k))
        return students, total, tariffs

    def starter(self, s, day):
        ws = next((w for w in self.v.worksheets if w.title.strip() == f"{s} спецкурсы"), None)
        if ws is None or not any(str(c.value).strip() == "Стартер-пак"
                                 for row in ws.iter_rows(min_row=2, max_col=5) for c in row if c.value):
            return None  # предмет стартер-пак не продаёт — строки в отчёте нет
        return sum(1 for r in range(2, ws.max_row + 1)
                   if str(ws.cell(r, 4).value or "").strip() == "Стартер-пак"
                   and (as_date(ws.cell(r, 1).value) or dt.date.max) <= day
                   and not (_fill(ws.cell(r, 1)) in REFUND_FILLS and _fill(ws.cell(r, 2)) in REFUND_FILLS))

    # ---------- текст ----------

    def report(self, day):
        plan = self.plans.get((day.year, day.month))
        if plan is None:
            raise ValueError(f"на листе «{ZP}» нет плана за {day:%m.%Y}")
        out = [f"📌 Отчет за {day.day} {GEN[day.month - 1]}:", "", "Всего обработано: "]
        for s in SUBJ:
            rows = self.logs[s]
            today = [a for _, d, a in rows if d == day]
            mtd = sum(a for _, d, a in rows if (d.year, d.month) == (day.year, day.month) and d <= day)
            # половину округляем вверх, как Excel/Sheets: 12.565 → 12.57 (float дал бы 12.56)
            pct = (Decimal(str(mtd)) * 100 / Decimal(str(plan[s]))).quantize(Decimal("0.01"), ROUND_HALF_UP) \
                if plan[s] else Decimal("0.00")
            students, total, tariffs = self.gk(s, day)
            out += ["", HEAD[s], "",
                    f"Получено оплат: {len(today)}",
                    f"Сумма оплат за сегодня: {money(sum(today))}",
                    f"Процент выполнения плана: {money(mtd) if round(mtd) else '0.00'} от {money(plan[s])} - {pct}%",
                    "",
                    f"Годовой курс: {students}",
                    f"В деньгах: {money(total)}",
                    " / ".join(f"{k} {v}" for k, v in tariffs)]
            sp = self.starter(s, day)
            if sp is not None:
                out += ["", f"Стартер-пак: {sp}"]
        return "\n".join(out)


def days(first, last=None):
    last = last or first
    if last < first:
        first, last = last, first
    return [first + dt.timedelta(n) for n in range((last - first).days + 1)]


def write_docx(texts, path):
    """Каждый отчёт — строка в строку как текст, каждый день с новой страницы."""
    doc = docx.Document()
    for i, text in enumerate(texts):
        if i:
            doc.paragraphs[-1].add_run().add_break(WD_BREAK.PAGE)
        for line in text.split("\n"):
            doc.add_paragraph(line)
    doc.save(path)
